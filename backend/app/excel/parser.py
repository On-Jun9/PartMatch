"""
엑셀 파일 파싱 및 데이터 추출 모듈
"""
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from typing import List, Dict, Any, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)

class ExcelParser:
    """거래명세서 엑셀 파일 파서"""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)
        self.workbook = None
        self.sheet: Optional[Worksheet] = None

    def load(self, sheet_name: str | int = 0):
        """엑셀 파일 로드"""
        try:
            self.workbook = load_workbook(self.file_path, data_only=True)

            # 시트 선택
            if isinstance(sheet_name, int):
                self.sheet = self.workbook.worksheets[sheet_name]
            else:
                self.sheet = self.workbook[sheet_name]

            logger.info(f"엑셀 파일 로드 완료: {self.file_path}")
            return self

        except Exception as e:
            logger.error(f"엑셀 파일 로드 실패: {e}")
            raise

    def get_cell_value(self, row: int, col: int) -> Any:
        """특정 셀 값 가져오기"""
        if not self.sheet:
            raise ValueError("먼저 load()를 호출하세요")
        return self.sheet.cell(row=row, column=col).value

    def analyze_structure(self) -> Dict[str, Any]:
        """엑셀 파일 구조 분석"""
        if not self.sheet:
            raise ValueError("먼저 load()를 호출하세요")

        # 기본 정보
        info = {
            "sheet_name": self.sheet.title,
            "max_row": self.sheet.max_row,
            "max_column": self.sheet.max_column,
            "dimensions": self.sheet.dimensions,
        }

        # 첫 10행 데이터 샘플
        sample_data = []
        for row in range(1, min(11, self.sheet.max_row + 1)):
            row_data = []
            for col in range(1, self.sheet.max_column + 1):
                value = self.sheet.cell(row=row, column=col).value
                row_data.append(value)
            sample_data.append(row_data)

        info["sample_data"] = sample_data

        return info

    def extract_items(
        self,
        start_row: int,
        item_name_col: int,
        quantity_col: int = None,
        unit_price_col: int = None,
        end_row: int = None
    ) -> List[Dict[str, Any]]:
        """
        거래명세서에서 품목 데이터 추출

        Args:
            start_row: 데이터 시작 행
            item_name_col: 품목명 열 번호
            quantity_col: 수량 열 번호
            unit_price_col: 단가 열 번호
            end_row: 데이터 종료 행 (None이면 마지막까지)
        """
        if not self.sheet:
            raise ValueError("먼저 load()를 호출하세요")

        items = []
        max_row = end_row or self.sheet.max_row

        for row in range(start_row, max_row + 1):
            item_name = self.sheet.cell(row=row, column=item_name_col).value

            # 품목명이 없으면 스킵
            if not item_name or str(item_name).strip() == "":
                continue

            item = {"name": str(item_name).strip()}

            if quantity_col:
                quantity = self.sheet.cell(row=row, column=quantity_col).value
                item["quantity"] = quantity

            if unit_price_col:
                unit_price = self.sheet.cell(row=row, column=unit_price_col).value
                item["unit_price"] = unit_price

            items.append(item)

        return items

    def close(self):
        """워크북 닫기"""
        if self.workbook:
            self.workbook.close()

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()


def analyze_excel_file(file_path: str | Path) -> Dict[str, Any]:
    """엑셀 파일 구조 분석 (편의 함수)"""
    with ExcelParser(file_path) as parser:
        parser.load()
        return parser.analyze_structure()
