"""
엑셀 거래명세서 생성 모듈
"""
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pathlib import Path
from datetime import datetime
from typing import List, Optional, Dict, Any
import logging
import shutil

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """거래명세서 엑셀 파일 생성기"""

    def __init__(self, template_path: Path):
        """
        Args:
            template_path: 템플릿 엑셀 파일 경로
        """
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(f"템플릿 파일을 찾을 수 없습니다: {template_path}")

    def generate(
        self,
        output_path: Path,
        items: List[Dict[str, Any]],
        date: Optional[datetime] = None,
        vehicle_number: Optional[str] = None,
        start_row: int = 14,
    ) -> Path:
        """
        거래명세서 엑셀 파일 생성

        Args:
            output_path: 출력 파일 경로
            items: 부품 아이템 리스트 (name, specification, quantity, unit_price, supply_price, tax_amount)
            date: 거래 날짜 (기본: 오늘)
            vehicle_number: 차 번호 (귀하) (기본: None)
            start_row: 품목 시작 행 (기본: 14)

        Returns:
            생성된 파일 경로
        """
        # 템플릿 파일 복사
        shutil.copy(self.template_path, output_path)
        logger.info(f"템플릿 파일 복사 완료: {self.template_path} -> {output_path}")

        # 워크북 열기
        workbook = load_workbook(output_path)
        sheet: Worksheet = workbook.active

        # 날짜 입력 (B6: 연도, F6: 월, I6: 일)
        if date is None:
            date = datetime.now()

        year = date.year
        month = date.month
        day = date.day

        sheet.cell(row=6, column=2).value = year  # B6: 연도
        sheet.cell(row=6, column=6).value = month  # F6: 월
        sheet.cell(row=6, column=9).value = day  # I6: 일

        logger.info(f"날짜 입력 완료: {year}년 {month}월 {day}일")

        # 차 번호(귀하) 입력 (B8)
        if vehicle_number:
            sheet.cell(row=8, column=2).value = vehicle_number
            logger.info(f"차 번호 입력 완료: {vehicle_number}")

        # 품목 데이터 입력 (14행부터)
        for idx, item in enumerate(items):
            row = start_row + idx

            # A열 (1): 품목명
            sheet.cell(row=row, column=1).value = item.get("name")

            # K열 (11): 규격
            if item.get("specification"):
                sheet.cell(row=row, column=11).value = item["specification"]

            # O열 (15): 수량
            if item.get("quantity") is not None:
                sheet.cell(row=row, column=15).value = item["quantity"]

            # R열 (18): 단가
            if item.get("unit_price") is not None:
                sheet.cell(row=row, column=18).value = item["unit_price"]

            # X열 (24): 공급가액
            supply_price = item.get("supply_price")
            if supply_price is None and item.get("quantity") and item.get("unit_price"):
                # 자동 계산: 수량 × 단가
                supply_price = item["quantity"] * item["unit_price"]
            if supply_price is not None:
                sheet.cell(row=row, column=24).value = supply_price

            # AH열 (34): 세액
            # 템플릿의 기존 수식을 유지하여 엑셀에서 자동 계산되도록 함
            # 사용자가 엑셀 다운로드 후 추가/수정 시 자동 계산 가능

            logger.debug(f"품목 입력: 행{row}, 품명='{item.get('name')}'")

        logger.info(f"총 {len(items)}개 품목 입력 완료")

        # 저장
        workbook.save(output_path)
        workbook.close()
        logger.info(f"엑셀 파일 생성 완료: {output_path}")

        return output_path
